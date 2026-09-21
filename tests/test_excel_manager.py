from pathlib import Path
from dataclasses import replace
import logging

import pandas as pd
from openpyxl import load_workbook

from excel_manager import (
    RETRY_REQUEST_STATUS,
    ExcelManager,
    ResultManager,
)
from main import _apply_retry_requests, _is_pending_record, _record_invalid_excel_rows


def _sample_frame() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "نام و نام خانوادگی نوآموز": "نمونه آزمایشی",
                "کد ملی نوآموز": 123456789,
                "تاریخ تولد نوآموز": 14000102,
                "محل تولد نوآموز": "بندرعباس",
                "نام پدر": "نام آزمایشی",
                "کد ملی پدر": 123456789,
                "تاریخ تولد پدر": 13600102,
                "مدرک تحصیلی پدر": "دیپلم",
                "شغل پدر": "آزاد",
                "نام و نام خانوادگی مادر": "نام آزمایشی",
                "کد ملی مادر": 987654321,
                "تاریخ تولد مادر": 13650102,
                "مدرک تحصیلی مادر": "لیسانس",
                "شغل مادر": None,
                "آدرس منزل": "نشانی آزمایشی",
                "کدپستی": 7464143583,
                "تلفن پدر": 9234567890,
                "تلفن مادر": 9123456789,
            }
        ]
    )


def _result_manager(source: Path, tmp_path: Path) -> ResultManager:
    return ResultManager(
        source,
        "اطلاعات نوآموزان",
        tmp_path / "results.completed.json",
    )


def test_load_and_normalize_identifiers(tmp_path: Path) -> None:
    source = tmp_path / "students.xlsx"
    _sample_frame().to_excel(source, sheet_name="اطلاعات نوآموزان", index=False)
    loaded = ExcelManager(source, "اطلاعات نوآموزان").load_students()
    assert loaded.issues == ()
    assert loaded.records[0].student_national_id == "0123456789"
    assert loaded.records[0].phone == "09123456789"


def test_missing_mother_phone_requires_review(tmp_path: Path) -> None:
    source = tmp_path / "students.xlsx"
    frame = _sample_frame()
    frame.loc[0, "تلفن مادر"] = None
    frame.to_excel(source, sheet_name="اطلاعات نوآموزان", index=False)

    loaded = ExcelManager(source, "اطلاعات نوآموزان").load_students()

    assert any(issue.field == "phone" for issue in loaded.issues)


def test_missing_postal_code_requires_review(tmp_path: Path) -> None:
    source = tmp_path / "students.xlsx"
    frame = _sample_frame()
    frame.loc[0, "کدپستی"] = None
    frame.to_excel(source, sheet_name="اطلاعات نوآموزان", index=False)

    loaded = ExcelManager(source, "اطلاعات نوآموزان").load_students()

    assert any(issue.field == "postal_code" for issue in loaded.issues)


def test_missing_father_phone_requires_review(tmp_path: Path) -> None:
    source = tmp_path / "students.xlsx"
    frame = _sample_frame()
    frame.loc[0, "تلفن پدر"] = None
    frame.to_excel(source, sheet_name="اطلاعات نوآموزان", index=False)

    loaded = ExcelManager(source, "اطلاعات نوآموزان").load_students()

    assert any(issue.field == "father_phone" for issue in loaded.issues)


def test_optional_manual_parent_columns_are_loaded_when_present(tmp_path: Path) -> None:
    source = tmp_path / "students.xlsx"
    frame = _sample_frame()
    frame["نام کوچک پدر"] = "علی"
    frame["نام خانوادگی پدر"] = "نمونه"
    frame["نام پدر پدر"] = "حسن"
    frame["نام کوچک مادر"] = "مریم"
    frame["نام خانوادگی مادر"] = "نمونه"
    frame["نام پدر مادر"] = "رضا"
    frame.to_excel(source, sheet_name="اطلاعات نوآموزان", index=False)

    record = ExcelManager(source, "اطلاعات نوآموزان").load_students().records[0]

    assert record.manual_father_info_complete
    assert record.manual_mother_info_complete


def test_robot_status_is_optional_and_loaded_when_present(tmp_path: Path) -> None:
    source = tmp_path / "students.xlsx"
    frame = _sample_frame()
    frame["وضعیت ربات"] = RETRY_REQUEST_STATUS
    frame.to_excel(source, sheet_name="اطلاعات نوآموزان", index=False)

    record = ExcelManager(source, "اطلاعات نوآموزان").load_students().records[0]

    assert record.robot_status == RETRY_REQUEST_STATUS


def test_results_are_upserted_and_success_is_resumable(tmp_path: Path) -> None:
    source = tmp_path / "students.xlsx"
    _sample_frame().to_excel(source, sheet_name="اطلاعات نوآموزان", index=False)
    manager = _result_manager(source, tmp_path)
    manager.upsert(
        row_number=2,
        student_national_id="0123456789",
        status="PENDING",
    )
    manager.upsert(
        row_number=2,
        student_national_id="0123456789",
        status="SUCCESS",
        selected_grade="پیش دبستانی ۱",
    )
    workbook = load_workbook(source, read_only=True, data_only=True)
    assert workbook.active.max_row == 2
    headers = [cell.value for cell in workbook.active[1]]
    assert headers[-4:] == [
        "وضعیت ربات",
        "پایه ثبت‌شده",
        "آخرین خطا",
        "زمان آخرین تلاش",
    ]
    assert workbook.active.cell(2, len(headers) - 2).value == "پیش دبستانی ۱"
    workbook.close()
    assert manager.successful_keys() == {(2, "0123456789")}
    assert manager.successful_national_ids() == {"0123456789"}


def test_success_recovery_survives_locked_students_workbook(
    tmp_path: Path,
    monkeypatch,
) -> None:
    source = tmp_path / "students.xlsx"
    _sample_frame().to_excel(source, sheet_name="اطلاعات نوآموزان", index=False)
    manager = _result_manager(source, tmp_path)

    def locked_upsert(**_kwargs) -> None:
        raise PermissionError("students.xlsx is open")

    monkeypatch.setattr(manager, "upsert", locked_upsert)

    assert not manager.record_success(
        row_number=2,
        student_national_id="0123456789",
        selected_grade="نوباوه",
    )
    assert manager.successful_national_ids() == {"0123456789"}


def test_recovery_file_syncs_to_students_workbook_on_next_run(tmp_path: Path) -> None:
    source = tmp_path / "students.xlsx"
    _sample_frame().to_excel(source, sheet_name="اطلاعات نوآموزان", index=False)
    manager = _result_manager(source, tmp_path)
    manager._record_completion(
        row_number=2,
        student_national_id="0123456789",
        selected_grade="نوباوه",
    )

    assert manager.successful_national_ids() == {"0123456789"}
    assert manager.successful_keys() == {(2, "0123456789")}


def test_successful_workbook_row_backfills_json_journal(tmp_path: Path) -> None:
    source = tmp_path / "students.xlsx"
    frame = _sample_frame()
    frame["وضعیت ربات"] = "ثبت شد"
    frame["پایه ثبت‌شده"] = "نوباوه"
    frame["زمان آخرین تلاش"] = "2026-09-20T12:00:00+03:30"
    frame.to_excel(source, sheet_name="اطلاعات نوآموزان", index=False)
    manager = _result_manager(source, tmp_path)

    manager.sync_completion_journal_from_workbook()

    assert manager._load_completion_records()["0123456789"]["selected_grade"] == (
        "نوباوه"
    )


def test_retry_request_overrides_previous_success(tmp_path: Path) -> None:
    source = tmp_path / "students.xlsx"
    frame = _sample_frame()
    frame["وضعیت ربات"] = RETRY_REQUEST_STATUS
    frame.to_excel(source, sheet_name="اطلاعات نوآموزان", index=False)
    loaded = ExcelManager(source, "اطلاعات نوآموزان").load_students()
    record = loaded.records[0]
    manager = _result_manager(source, tmp_path)
    manager.record_success(
        row_number=record.source_row,
        student_national_id=record.student_national_id or "",
        selected_grade="پیش دبستانی ۱",
    )

    _apply_retry_requests(loaded, manager, logging.getLogger("test"))

    assert manager.successful_national_ids() == set()
    assert manager.latest_statuses()[record.student_national_id or ""] == "PENDING"


def test_result_columns_are_created_and_updated_in_students_workbook(
    tmp_path: Path,
) -> None:
    source = tmp_path / "students.xlsx"
    _sample_frame().to_excel(source, sheet_name="اطلاعات نوآموزان", index=False)
    loaded = ExcelManager(source, "اطلاعات نوآموزان").load_students()
    record = loaded.records[0]
    results = _result_manager(source, tmp_path)
    results.upsert(
        row_number=record.source_row,
        student_national_id=record.student_national_id or "",
        status="INCOMPLETE_FATHER_INFO",
    )

    workbook = load_workbook(source, read_only=True, data_only=True)
    worksheet = workbook["اطلاعات نوآموزان"]
    headers = [cell.value for cell in worksheet[1]]
    assert headers[-4:] == [
        "وضعیت ربات",
        "پایه ثبت‌شده",
        "آخرین خطا",
        "زمان آخرین تلاش",
    ]
    assert worksheet.cell(2, len(headers) - 3).value == (
        "بررسی: اطلاعات پدر ناقص است"
    )
    workbook.close()


def test_parent_inquiry_failure_has_a_distinct_result_status(tmp_path: Path) -> None:
    source = tmp_path / "students.xlsx"
    _sample_frame().to_excel(source, sheet_name="اطلاعات نوآموزان", index=False)
    manager = _result_manager(source, tmp_path)

    manager.upsert(
        row_number=2,
        student_national_id="0123456789",
        status="INCOMPLETE_PARENT_INFO",
        error_message="اطلاعات مادر ناقص است",
    )

    workbook = load_workbook(source, read_only=True, data_only=True)
    worksheet = workbook.active
    headers = {cell.value: cell.column for cell in worksheet[1]}
    status_value = worksheet.cell(2, headers["وضعیت ربات"]).value
    error_value = worksheet.cell(2, headers["آخرین خطا"]).value
    workbook.close()
    assert status_value == "بررسی: اطلاعات والدین ناقص است"
    assert error_value == "اطلاعات مادر ناقص است"


def test_incomplete_parent_reenters_queue_only_after_manual_fields_are_complete(
    tmp_path: Path,
) -> None:
    source = tmp_path / "students.xlsx"
    _sample_frame().to_excel(source, sheet_name="اطلاعات نوآموزان", index=False)
    record = ExcelManager(source, "اطلاعات نوآموزان").load_students().records[0]
    national_id = record.student_national_id or ""

    assert not _is_pending_record(
        record,
        set(),
        {national_id: "INCOMPLETE_FATHER_INFO"},
    )
    completed = replace(
        record,
        manual_father_first_name="علی",
        manual_father_last_name="نمونه",
        manual_father_parent_name="حسن",
    )
    assert _is_pending_record(
        completed,
        set(),
        {national_id: "INCOMPLETE_FATHER_INFO"},
    )
    assert not _is_pending_record(
        completed,
        set(),
        {national_id: "INCOMPLETE_STUDENT_INFO"},
    )
    assert not _is_pending_record(
        completed,
        set(),
        {national_id: "INVALID_STUDENT_NATIONAL_ID"},
    )


def test_invalid_excel_row_is_written_to_results_and_held(tmp_path: Path) -> None:
    source = tmp_path / "students.xlsx"
    frame = _sample_frame()
    frame.loc[0, "تاریخ تولد پدر"] = None
    frame.to_excel(source, sheet_name="اطلاعات نوآموزان", index=False)
    loaded = ExcelManager(source, "اطلاعات نوآموزان").load_students()
    manager = _result_manager(source, tmp_path)

    _record_invalid_excel_rows(loaded, manager, logging.getLogger("test"))

    record = loaded.records[0]
    national_id = record.student_national_id or ""
    statuses = manager.latest_statuses()
    assert statuses[national_id] == "INVALID_EXCEL_DATA"
    assert not _is_pending_record(record, set(), statuses)
    workbook = load_workbook(source, read_only=True, data_only=True)
    worksheet = workbook.active
    headers = {cell.value: cell.column for cell in worksheet[1]}
    last_error = worksheet.cell(2, headers["آخرین خطا"]).value
    workbook.close()
    assert "father_birth_date" in last_error
